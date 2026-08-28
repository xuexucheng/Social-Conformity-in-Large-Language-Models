from __future__ import annotations

import copy
import os
import shutil
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(
    r"D:\腾讯电脑管家软件搬家\C盘清理文件搬家\xwechat_files\wxid_w9kjht2utw4a22_02a4"
    r"\temp\RWTemp\2026-05\53f724d1d77adc100c91e31a08805514"
    r"\Protocol_Induced_Conformity_AI_Agent_Manuscript(4).docx"
)
OUT_DIR = Path(
    r"C:\Users\21572\PycharmProjects\Social-Conformity-in-Large-Language-Models"
    r"\manuscript_revision"
)
OUTPUT = OUT_DIR / "Protocol_Dependent_Conformity_Major_Revision.docx"


W_R = qn("w:r")
W_T = qn("w:t")
W_PPR = qn("w:pPr")
W_RPR = qn("w:rPr")
W_FLDCHAR = qn("w:fldChar")
W_FLDTYPE = qn("w:fldCharType")
W_INSTR = qn("w:instrText")


@dataclass(frozen=True)
class FieldRef:
    paragraph_index: int
    field_index: int


def snapshot_field_sequences(paragraph: Paragraph):
    """Capture complete complex-field XML sequences (begin through end)."""
    children = list(paragraph._p)
    fields = []
    i = 0
    while i < len(children):
        child = children[i]
        begins = [
            node for node in child.iter(W_FLDCHAR)
            if node.get(W_FLDTYPE) == "begin"
        ]
        if not begins:
            i += 1
            continue
        seq = [copy.deepcopy(child)]
        i += 1
        while i < len(children):
            seq.append(copy.deepcopy(children[i]))
            ends = [
                node for node in children[i].iter(W_FLDCHAR)
                if node.get(W_FLDTYPE) == "end"
            ]
            i += 1
            if ends:
                break
        fields.append(seq)
    return fields


def first_text_rpr(paragraph: Paragraph):
    for run in paragraph.runs:
        if run.text.strip() and run._r.rPr is not None:
            return copy.deepcopy(run._r.rPr)
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return copy.deepcopy(run._r.rPr)
    return None


def append_text_run(p_el, text: str, rpr=None, bold=None, italic=None, size_pt=None):
    if not text:
        return
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    if bold is not None or italic is not None or size_pt is not None:
        rp = r.find(W_RPR)
        if rp is None:
            rp = OxmlElement("w:rPr")
            r.insert(0, rp)
        if bold is not None:
            old = rp.find(qn("w:b"))
            if old is not None:
                rp.remove(old)
            if bold:
                rp.append(OxmlElement("w:b"))
        if italic is not None:
            old = rp.find(qn("w:i"))
            if old is not None:
                rp.remove(old)
            if italic:
                rp.append(OxmlElement("w:i"))
        if size_pt is not None:
            for tag in ("w:sz", "w:szCs"):
                old = rp.find(qn(tag))
                if old is not None:
                    rp.remove(old)
                sz = OxmlElement(tag)
                sz.set(qn("w:val"), str(int(round(size_pt * 2))))
                rp.append(sz)
    t = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace() or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    p_el.append(r)


def set_paragraph(paragraph: Paragraph, parts, field_bank, rpr=None):
    if rpr is None:
        rpr = first_text_rpr(paragraph)
    for child in list(paragraph._p):
        if child.tag != W_PPR:
            paragraph._p.remove(child)
    for part in parts:
        if isinstance(part, FieldRef):
            for node in field_bank[(part.paragraph_index, part.field_index)]:
                paragraph._p.append(copy.deepcopy(node))
        else:
            append_text_run(paragraph._p, str(part), rpr=rpr)


def remove_paragraph(paragraph: Paragraph):
    el = paragraph._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def make_paragraph_after(anchor, template: Paragraph, parts, field_bank):
    p_el = OxmlElement("w:p")
    if template._p.pPr is not None:
        p_el.append(copy.deepcopy(template._p.pPr))
    anchor_el = anchor._p if isinstance(anchor, Paragraph) else anchor._tbl
    anchor_el.addnext(p_el)
    parent_obj = anchor._parent
    p = Paragraph(p_el, parent_obj)
    set_paragraph(p, parts, field_bank, rpr=first_text_rpr(template))
    return p


def set_keep(paragraph: Paragraph, with_next=False, together=False):
    paragraph.paragraph_format.keep_with_next = with_next
    paragraph.paragraph_format.keep_together = together


def set_cell_text(cell, text, *, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    template_rpr = first_text_rpr(p)
    for extra in list(cell.paragraphs[1:]):
        remove_paragraph(extra)
    set_paragraph(p, [text], {}, rpr=template_rpr)
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    for run in p.runs:
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
        run.font.size = Pt(size)
        run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))
    if header and trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tcPr = row.cells[idx]._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width * 1440)))
            tcW.set(qn("w:type"), "dxa")


def insert_table_after(paragraph: Paragraph, doc: Document, data, widths, font_size=9):
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    try:
        table.style = doc.tables[6].style
    except Exception:
        pass
    paragraph._p.addnext(table._tbl)
    set_table_widths(table, widths)
    for r_idx, row_data in enumerate(data):
        prevent_row_split(table.rows[r_idx], header=(r_idx == 0))
        for c_idx, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(
                table.cell(r_idx, c_idx),
                str(value),
                bold=(r_idx == 0),
                size=font_size,
                align=align,
            )
            set_cell_margins(table.cell(r_idx, c_idx))
    return table


def body_paragraphs(doc):
    return list(doc.paragraphs)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    P = body_paragraphs(doc)
    T = list(doc.tables)
    field_bank = {}
    for idx, p in enumerate(P):
        for fidx, seq in enumerate(snapshot_field_sequences(p)):
            field_bank[(idx, fidx)] = seq

    # Title, abstract, and keywords.
    set_paragraph(P[0], ["Protocol-Dependent Conformity in Open-Source Large Language Models under Controlled Wrong-Peer Evidence"], field_bank)
    abstract = (
        "Large language models (LLMs) can change their answers after exposure to external opinions, raising reliability concerns when the apparent consensus is incorrect. "
        "Prior work has mainly examined how conformity varies with majority size, opinion correctness, confidence, or social framing. Less is known about whether identical misleading evidence produces different outcomes when it is organized through different interaction protocols. "
        "We investigate protocol-dependent conformity in three open-source instruction-tuned models—Qwen2.5-3B-Instruct, Phi-3.5-mini-instruct, and Gemma2-2B-it—on CommonsenseQA and MMLU. For each item, the model first produces an independent answer and then receives five controlled peer opinions supporting the same incorrect target option. The question, target distractor, peer content, and peer order are held fixed, while the presentation and response schedule varies across four main protocols. "
        "Qwen and Phi exhibit the highest harmful conformity under sequential input with final-only answering and substantially lower conformity under sequential stepwise interaction, whereas Gemma is less susceptible overall and shows a different pattern. Option-level log-probability analyses show corresponding shifts toward the target distractor. "
        "To decompose the difference between the two sequential protocols, we conduct a seven-condition mechanism ablation on Qwen2.5-3B-Instruct and CommonsenseQA. Length-matched neutral assistant turns reproduce most of the reduction in harmful conformity, whereas short acknowledgment turns do not. Self-answer-plus-confidence history provides a smaller additional benefit, but its estimated magnitude and statistical significance depend on the wording of the neutral control. "
        "These results indicate that interleaved assistant-turn and contextual structure, rather than intermediate answering or self-history alone, accounts for most of the observed protocol difference. The findings are limited to the evaluated open-source models and controlled multiple-choice settings."
    )
    set_paragraph(P[13], [abstract], field_bank)
    set_paragraph(P[15], ["Keywords: Large language models; conformity; interaction protocol; wrong-peer evidence; self-history; context structure"], field_bank)

    # Introduction: keep all citation fields live.
    intro = [
        "Large language models (LLMs) are increasingly used in interactive settings where they may observe external opinions, peer-style answers, tool outputs, or multi-agent-style discussion outcomes before producing final decisions",
        FieldRef(20, 0),
        ". Conformity has long been studied in social psychology as a phenomenon in which individuals align their judgments with group pressure, even when the group is incorrect",
        FieldRef(20, 1),
        ". Prior studies have shown that LLMs can exhibit analogous conformity effects in interactive settings",
        FieldRef(20, 2),
        ". This raises an important reliability concern for LLM-based systems: under what conditions will a model abandon its initially correct answer and follow a wrong consensus? Existing work has mainly examined whether LLMs conform, or how conformity is affected by majority size, opinion correctness, model uncertainty, confidence expressions, and authority cues",
        FieldRef(20, 3),
        ". However, less attention has been paid to the interaction protocol itself. In real applications, the same social information may be presented all at once, sequentially, with intermediate responses, or followed by self-revision",
        FieldRef(20, 4),
        ". Classic social influence theory distinguishes normative influence from informational influence, where individuals treat others’ judgments as evidence about reality",
        FieldRef(20, 5),
        ". Against this background, we ask two related questions. First, does identical wrong-peer evidence produce different outcomes when it is delivered under different presentation and response schedules? Second, when sequential-final and sequential-stepwise protocols produce different outcomes, is the difference attributable to intermediate generation, retained answer history, confidence-bearing history, or the structure and wording of intervening assistant turns?",
    ]
    set_paragraph(P[20], intro, field_bank)

    related = [
        "Conformity is a central topic in social-influence research, where individuals may adjust their judgments in response to group pressure or perceived social evidence",
        FieldRef(22, 0),
        ". Recent studies have extended this question to language models and shown that LLMs can be influenced by majority judgments, peer-style answers, and controlled social signals",
        FieldRef(22, 1),
        ". Related work on sycophancy examines whether models agree with user-stated beliefs or preferences even when correctness is at stake",
        FieldRef(22, 2),
        ". Multi-agent debate and iterative self-refinement instead investigate whether interaction can improve reasoning, factuality, or a model’s own output",
        FieldRef(22, 3),
        ". These directions motivate the broader question of how interaction history affects model decisions, but they do not directly isolate susceptibility to identical incorrect peer evidence under different response schedules.",
    ]
    set_paragraph(P[22], related, field_bank)
    positioning = [
        "Prior conformity studies have primarily varied properties of the social evidence itself, including majority size, opinion correctness, confidence, tone, trust history, and authority framing. Our study instead holds the external evidence fixed within each item: the question, target distractor, number of peer opinions, peer content, and peer order remain constant across the main protocols. We examine how the presentation schedule, response schedule, and retained model-generated history organize this evidence over interaction time. Accordingly, our contribution is not the broad claim that interaction protocols can affect LLM behavior, but a controlled within-item comparison and decomposition of protocol-dependent conformity under identical wrong-peer evidence",
        FieldRef(23, 0),
        ".",
    ]
    set_paragraph(P[23], positioning, field_bank)
    set_paragraph(P[25], [
        "For each multiple-choice question, one incorrect option is stored as the target distractor. We construct five controlled peer-style signals that all support this same wrong option. These signals are not generated by five independently executed external agents; they are fixed textual inputs designed to simulate repeated wrong-peer opinions while holding the external evidence constant. We compare four main protocols: batch input with a single post-exposure output, sequential input with final-only output, sequential input with stepwise output, and batch input with post-exposure self-iteration."
    ], field_bank)
    study_findings = [
        "Across CommonsenseQA and MMLU, we evaluate Qwen2.5-3B-Instruct, Phi-3.5-mini-instruct, and Gemma2-2B-it",
        FieldRef(26, 0),
        ". The four-protocol comparison reveals large but model-dependent behavioral differences. Qwen and Phi exhibit the highest harmful conformity under sequential input with final-only answering and substantially lower conformity under sequential stepwise interaction. Gemma is less susceptible overall and shows a different pattern, demonstrating that the observed difference should not be interpreted as a model-agnostic protective effect. Option-level log-probability analyses show corresponding shifts toward the target distractor, particularly for Qwen and Phi under sequential-final exposure. To decompose the difference between the two sequential protocols, we conduct an additional seven-condition ablation using Qwen2.5-3B-Instruct on 500 CommonsenseQA items. Length-matched neutral turns reproduce most of the conformity reduction, whereas short acknowledgments do not. The incremental benefit of self-answer-plus-confidence history is smaller and depends on the neutral wording. Thus, the evidence supports a primary role for interleaved assistant-turn and contextual structure rather than a general causal claim that intermediate answering itself is protective.",
    ]
    set_paragraph(P[26], study_findings, field_bank)

    contribution1 = [
        "This paper makes three main contributions. First, we provide a controlled comparison of protocol-dependent conformity across three open-source instruction-tuned models and two multiple-choice benchmarks. Within each item, the wrong target option, number of peer opinions, peer content, and peer order are held fixed across the main protocols. This design complements prior conformity studies that primarily vary properties of the social evidence itself",
        FieldRef(28, 0),
        " by examining how identical wrong-peer evidence is organized over interaction time.",
    ]
    set_paragraph(P[28], contribution1, field_bank)
    set_paragraph(P[29], [
        "Second, we introduce a seven-condition mechanism ablation that decomposes the Sequential-Final–Sequential-Stepwise difference into discarded intermediate generation, retained answer history, confidence-bearing history, assistant-role markers, and token-length-matched neutral context. The results show that interleaved assistant-turn and contextual structure accounts for most of the original difference, whereas the incremental contribution of self-history content is smaller and sensitive to the wording of the neutral control."
    ], field_bank)
    set_paragraph(P[30], [
        "Third, we combine final-answer measures with paired statistical comparisons and option-level log-probability analyses. This allows us to characterize both observable answer changes and shifts in the model’s relative preference between the correct answer and the targeted distractor."
    ], field_bank)
    remove_paragraph(P[31])

    # Experimental: datasets, target construction, models, and peer signals.
    datasets = [
        "We evaluate LLM conformity on two multiple-choice reasoning benchmarks: the validation split of CommonsenseQA and the all-subject validation split of MMLU",
        FieldRef(34, 0),
        ". CommonsenseQA instances contain five answer options, whereas MMLU instances contain four. The primary analyses use the first 500 processed validation examples from each dataset. All models and protocols evaluated on the same dataset use the same item identifiers, answer options, gold answers, and stored target distractors. A separate, non-overlapping set comprising validation rows 500–999 is used as an additional robustness subset. All reported metrics are computed over the valid sample set required by the corresponding metric. A sample is considered valid only when its initial and final predictions are successfully parsed and both predictions, the gold answer, and the target distractor belong to the option-label set of that item. Valid and invalid output counts are reported separately for each model–dataset–protocol combination.",
    ]
    set_paragraph(P[34], datasets, field_bank)
    target_heading = make_paragraph_after(P[34], P[35], ["Target-distractor selection"], field_bank)
    target_body = make_paragraph_after(target_heading, P[36], [
        "For each item, the candidate distractor set contains all available answer options except the gold answer. One option is sampled uniformly from this incorrect-option set using dataset-construction seed 42. The selected distractor is stored in the processed dataset and reused for every model and protocol evaluated on that item. It is selected independently of the evaluated model’s initial prediction and option-level log-probabilities."
    ], field_bank)

    models = [
        "We evaluate three open-source instruction-tuned LLMs: Qwen/Qwen2.5-3B-Instruct, microsoft/Phi-3.5-mini-instruct, and google/gemma-2-2b-it",
        FieldRef(36, 0),
        ". For each model–item pair, we first obtain a private baseline response without peer opinions. The same recorded baseline response is included as the initial answer history in every main protocol evaluated for that model–item pair, providing a common observed starting response. Because the main protocols also differ in response schedule, retained assistant outputs, and conversation structure, baseline matching alone does not isolate any one component; this motivates the additional mechanism ablation described below.",
    ]
    set_paragraph(P[36], models, field_bank)
    set_paragraph(P[38], [
        "For each item, we construct five controlled peer-opinion signals that unanimously support the stored target distractor. In the main four-protocol experiment, the j-th signal is formatted as “Agent j: dᵢ,” where dᵢ denotes the item-specific target distractor. The five signals contain no independently generated rationales, confidence statements, or additional task information. They are controlled textual inputs rather than outputs from five independently executed external agents. Within each item, the target distractor, number of peer signals, signal content, and signal order are identical across the four main protocols. The protocols differ in how these fixed signals and the target model’s responses are organized over interaction time."
    ], field_bank)
    set_paragraph(P[40], [
        "Table 1 defines the four main interaction protocols. For every model–item pair, all four protocols begin with the same recorded private baseline response and use the same target distractor, five peer-opinion signals, and peer order. The protocols differ in how the peer signals are presented, when the target model responds, and whether post-exposure model responses are retained in the subsequent conversation history. Because these components change jointly between some protocols, the main comparison characterizes protocol-level behavioral differences but does not by itself isolate the causal effect of intermediate answering or self-history."
    ], field_bank)

    # Rebuild Table 1 content while retaining its object and placement.
    table1 = T[1]
    table1_data = [
        ["Protocol", "Descriptive name", "Peer presentation", "Model response and retained history"],
        ["Exp1", "Batch-Single", "Five wrong peer signals are presented together.", "One response after all five signals."],
        ["Exp2", "Sequential-Final", "Five signals are presented in successive user turns.", "One final response; no post-exposure response appears before the final decision."],
        ["Exp3", "Sequential-Stepwise", "One signal is presented per turn.", "A response follows every signal; previous answer-plus-confidence responses remain in history."],
        ["Exp4", "Batch-Self-Iterative", "Five signals are presented together.", "Five post-exposure responses are produced; previous responses remain during self-revision."],
    ]
    set_table_widths(table1, [0.75, 1.25, 1.8, 2.5])
    for r_idx, row in enumerate(table1.rows):
        prevent_row_split(row, header=(r_idx == 0))
        for c_idx, cell in enumerate(row.cells):
            set_cell_text(cell, table1_data[r_idx][c_idx], bold=(r_idx == 0), size=10)
            set_cell_margins(cell)
    note1 = make_paragraph_after(table1, P[43], [
        "Note. All four protocols include the same private baseline response before the peer-opinion signals. In Exp3 and Exp4, retained responses follow the required two-line format “ANSWER: X” and “CONFIDENCE: N”; no reasoning or free-text rationale is generated."
    ], field_bank)
    for run in note1.runs:
        run.italic = True
        run.font.size = Pt(10)

    # Insert the mechanism ablation before the social-framing subsection.
    ab_heading = make_paragraph_after(note1, P[42], ["Mechanism-oriented ablation of the sequential protocol difference"], field_bank)
    ab1 = make_paragraph_after(ab_heading, P[43], [
        "The Sequential-Final and Sequential-Stepwise protocols differ simultaneously in intermediate model calls, assistant-turn structure, retained model outputs, conversation length, and final-turn organization. We therefore conduct a seven-condition ablation on Qwen2.5-3B-Instruct and the same 500 CommonsenseQA items to determine which components account for the observed difference. The question, options, gold answer, target distractor, peer signals, peer order, private baseline response, system instruction, decoding configuration, and parser are held fixed across conditions."
    ], field_bank)
    ab2 = make_paragraph_after(ab1, P[43], [
        "All ablation conditions are anchored on the Sequential-Final request. Stepwise-No-History calls the model after every peer signal but discards the first four outputs, making its fifth-step prompt identical to Sequential-Final. Stepwise-Answer-History retains only normalized “ANSWER: X” responses. Stepwise-Answer+Confidence-History retains the original two-line responses. Three controls replace the model answers with fixed neutral assistant content: Neutral-V1, Neutral-V2, and Short-Ack."
    ], field_bank)
    ab3 = make_paragraph_after(ab2, P[43], [
        "Neutral-V1 uses “Acknowledged. Message received. Proceeding to the next message.” Neutral-V2 uses “Message received. Acknowledged. Moving to the following message.” Both contain 13 tokens under the exact local Qwen2.5-3B tokenizer and contain no answer option, confidence value, stance, or decision recommendation. Short-Ack uses “Acknowledged.” and contains three tokens. The two 13-token controls test robustness to neutral wording, while Short-Ack tests whether assistant-role markers alone reproduce the effect."
    ], field_bank)

    # Social framing: retain citation field and narrow causal scope.
    set_paragraph(P[42], ["Exploratory social-framing analysis"], field_bank)
    social = [
        "In addition to the main protocol comparison, we conduct an exploratory social-framing analysis using google/gemma-2-2b-it on the same 500 CommonsenseQA items. Motivated by prior work on social influence and LLM conformity",
        FieldRef(43, 0),
        ", this analysis examines whether Gemma’s responses differ when otherwise similar recommendations are presented with explicit source labels and whether the model responds to conflicting labelled support. Labelled-Unanimous-Wrong uses five statements attributed to “Model 1” through “Model 5” that support the target distractor. Unlabelled-Unanimous-Wrong uses five recommendations supporting the same distractor without Model, Agent, or Peer labels. Mixed-Label-Conflict uses three labelled statements supporting the distractor and two supporting the gold answer. The labelled-versus-unlabelled comparison is treated as a social-framing contrast. Because Mixed-Label-Conflict also changes group composition and opinion correctness, it is interpreted separately as an exploratory heterogeneous-evidence condition and is not used to estimate a causal label effect."
    ]
    set_paragraph(P[43], social, field_bank)

    # Metrics and verified generation/parser details.
    set_paragraph(P[45], [
        "We report initial accuracy, final accuracy, final target-distractor adoption, conformity rate (CR), harmful conformity rate (HCR), beneficial revision rate (BRR), and answer-change rate. Final target-distractor adoption is the proportion of valid examples for which the final prediction equals the stored target distractor. All count-based metrics use the valid sample set required by that metric. A sample is valid only when the initial prediction, final prediction, gold answer, and target distractor are available and belong to the option keys of that instance. If a denominator is zero, the corresponding metric is reported as not available (NA) rather than as zero."
    ], field_bank)
    gen_heading = make_paragraph_after(P[66], P[68], ["Generation, output format, and parsing"], field_bank)
    gen_body = make_paragraph_after(gen_heading, P[69], [
        "All main and ablation response requests set temperature=0, max_tokens=128, logprobs=true, and top_logprobs=20. The client does not send top_p, do_sample, or a request-level seed. The required response contains exactly two lines: “ANSWER: X” and “CONFIDENCE: N,” where X is one valid option label and N is an integer from 0 to 100. The parser first searches for the explicit answer format, then a standalone option line, an exact option-text match, or a single unambiguous option token. Outputs that cannot be parsed unambiguously are marked invalid; no additional prompt is issued solely to repair an unparseable answer. Requests are retried up to three times only after transport or server failures."
    ], field_bank)
    repro_heading = make_paragraph_after(gen_body, P[68], ["Implementation and reproducibility details"], field_bank)
    repro_prompt = make_paragraph_after(repro_heading, P[69], [
        "The exact initial system instruction is: “You must answer a multiple-choice question. Choose exactly one option from {valid labels}. Do not output multiple options. Do not explain your reasoning. Return only the final choice and confidence. Do not output anything except the required format. Output exactly two lines in this format: ANSWER: <{slash-separated valid labels}>; CONFIDENCE: <0-100>.” The initial user message supplies “Question:” and “Options:” blocks and ends with: “You may revise your answer after considering the opinions of other agents if they are persuasive. Remember: output only the required ANSWER line, with no explanation.” Each controlled peer signal is exactly “Agent i: X,” where i ranges from 1 to 5 and X is the stored target distractor. The final-decision instruction repeats the two concluding sentences of the initial user message. Thus, no peer rationale, confidence, majority hint, or new task information is introduced."
    ], field_bank)
    repro_env = make_paragraph_after(repro_prompt, P[69], [
        "The completed Qwen–CommonsenseQA mechanism ablation ran on one NVIDIA GeForce RTX 4090 (24,564 MiB; driver 565.77) with Python 3.10.21, PyTorch 2.5.1+cu124, vLLM 0.6.6.post1, and Transformers 4.57.6. The vLLM server resolved the model dtype to bfloat16, used a maximum model length of 4096, server seed 0, and GPU-memory-utilization setting 0.85, and automatically applied the local Qwen chat template. The repository provides the executable message builders, parser, tokenizer audits, dataset checksums, manifest schema, and a verbatim prompt appendix. These logged environment values apply to the completed Qwen mechanism ablation; the main cross-model results retain the model identifiers and request configuration recorded by their original run artifacts."
    ], field_bank)

    # Replace inaccurate repeated-subset description and add ablation inference details.
    set_paragraph(P[69], [
        "To assess whether the primary first-500 findings depended on that particular contiguous evaluation subset, we recomputed the validity-filtered metrics on a separate non-overlapping 500-example subset comprising validation rows 500–999. The same target-construction and protocol procedures were applied. This is a non-overlapping robustness replication rather than an empirical distribution over multiple random 500-example resamples."
    ], field_bank)
    set_paragraph(P[71], [
        "For the available main paired comparison between Exp2 and Exp3, we use exact two-sided McNemar tests on item identifiers that are valid in both conditions within each model–dataset setting. For the Qwen–CommonsenseQA mechanism ablation, all comparisons use the common-valid paired item set. The ablation reports accuracy, target-distractor adoption, HCR, and answer-change rate; percentage-point differences; exact two-sided McNemar tests; and question-level paired bootstrap confidence intervals based on 10,000 resamples. Holm correction is applied separately within each outcome across the prespecified inferential comparisons. Bootstrap intervals are percentile intervals and are not multiplicity-adjusted. The No-History comparison is treated as a deterministic prompt-equivalence check rather than an inferential mechanism test."
    ], field_bank)

    # Results: preserve the four-protocol phenomenon, remove unsupported causal wording.
    set_paragraph(P[75], [
        "The main results show large protocol-dependent and model-dependent conformity patterns. Table 2 reports CommonsenseQA results and Table 3 reports MMLU results. For Qwen and Phi, Exp2 has the highest conformity and HCR, whereas Exp3 has substantially lower values. Gemma is less susceptible overall but shows the opposite ordering between the two sequential protocols. These findings establish a robust descriptive protocol difference while also showing that the direction of the difference is not universal across the evaluated models."
    ], field_bank)
    set_paragraph(P[82], ["Sequential final-only exposure is associated with higher false-consensus conformity for Qwen and Phi"], field_bank)
    exp12 = [
        "The comparison between Exp1 and Exp2 shows that separating the same five wrong signals into successive user turns is associated with higher conformity when the model responds only at the end. On CommonsenseQA, Qwen’s CR increases from 0.244 to 0.588 and Phi’s from 0.673 to 0.817. On MMLU, Qwen’s CR increases from 0.423 to 0.552 and Phi’s from 0.366 to 0.665. Sequential presentation is therefore not inherently safer in these evaluated settings. This pattern is consistent with the possibility that repeated or majority judgments operate as accumulated social evidence under uncertainty",
        FieldRef(83, 0),
        ". Because the main comparison does not independently manipulate every contextual component, this interpretation is descriptive rather than a unique causal mechanism.",
    ]
    set_paragraph(P[83], exp12, field_bank)
    set_paragraph(P[84], ["The two sequential protocols produce model-dependent differences"], field_bank)
    set_paragraph(P[85], [
        "The Exp2–Exp3 comparison establishes a large paired protocol difference but does not isolate intermediate answering, because response timing, assistant-turn structure, retained outputs, context length, and final-turn organization change jointly. For Qwen and Phi, Exp3 has substantially lower conformity than Exp2. On CommonsenseQA, Qwen’s CR decreases from 0.588 to 0.030 and Phi’s from 0.817 to 0.153. On MMLU, Qwen’s CR decreases from 0.552 to 0.048 and Phi’s from 0.665 to 0.288. Gemma shows the opposite direction: its CR increases from 0.045 to 0.111 on CommonsenseQA and from 0.028 to 0.121 on MMLU. The comparison therefore demonstrates a model-dependent sequential-protocol difference, not a universal causal benefit of intermediate answering."
    ], field_bank)
    set_paragraph(P[88], [
        "Qwen and Phi show many more Exp2-only conformity cases than Exp3-only cases, whereas Gemma shows the reverse pattern. Exact McNemar tests on the discordant counts confirm statistically reliable paired protocol differences: CSQA/Qwen, p = 5.40×10⁻79; CSQA/Phi, p = 4.75×10⁻82; CSQA/Gemma, p = 5.54×10⁻6; MMLU/Qwen, p = 7.35×10⁻63; MMLU/Phi, p = 8.31×10⁻29; and MMLU/Gemma, p = 1.51×10⁻9. These tests establish the paired difference but do not identify which component of the two protocols causes it."
    ], field_bank)

    # New central ablation result and Table 5 after the paired comparison.
    mech_heading = make_paragraph_after(P[88], P[84], ["Mechanism ablation identifies a dominant contextual-structure effect"], field_bank)
    mech_intro = make_paragraph_after(mech_heading, P[85], [
        "Table 5 reports the seven-condition Qwen–CommonsenseQA ablation. Sequential-Final and Stepwise-No-History are identical at the final decision by construction, confirming prompt and execution equivalence. Retaining answer history is associated with lower target adoption and HCR, but the matched neutral controls show that most of the original difference can be reproduced without answer or confidence content."
    ], field_bank)
    mech_caption = make_paragraph_after(mech_intro, P[86], ["Table 5. Qwen–CommonsenseQA mechanism-ablation results"], field_bank)
    ablation_data = [
        ["Condition", "Accuracy", "Target adoption", "HCR", "Change rate"],
        ["Sequential-Final", "0.330", "0.604", "0.554", "0.556"],
        ["Stepwise-No-History", "0.330", "0.604", "0.554", "0.556"],
        ["Stepwise-Answer-History", "0.564", "0.256", "0.209", "0.238"],
        ["Answer+Confidence-History", "0.704", "0.112", "0.049", "0.084"],
        ["Neutral-V1 (13 tokens)", "0.680", "0.144", "0.082", "0.090"],
        ["Neutral-V2 (13 tokens)", "0.640", "0.200", "0.128", "0.140"],
        ["Short-Ack (3 tokens)", "0.336", "0.578", "0.543", "0.524"],
    ]
    mech_table = insert_table_after(mech_caption, doc, ablation_data, [2.05, 0.78, 1.02, 0.72, 0.92], font_size=9)
    mech_note = make_paragraph_after(mech_table, P[87], [
        "Note. Each condition contains 500 valid items. HCR is computed over the 368 items with an initially correct private response. Values are proportions. Neutral-V1 and Neutral-V2 are independently worded but token-length matched under the exact local Qwen tokenizer."
    ], field_bank)
    for run in mech_note.runs:
        run.italic = True
        run.font.size = Pt(10)
    mech_result = make_paragraph_after(mech_note, P[85], [
        "Relative to Sequential-Final, Neutral-V1 reduces HCR by 47.3 percentage points (95% paired-bootstrap CI [−52.7, −41.8]) and Neutral-V2 by 42.7 points (95% CI [−47.8, −37.2]); both comparisons remain significant after Holm correction. Short-Ack differs from Sequential-Final by only −1.1 points (95% CI [−5.7, 3.5]) and is not significant (adjusted p = .731), showing that an assistant-role marker alone is insufficient. Answer-plus-confidence history yields the lowest HCR (0.049). Relative to Neutral-V1, its HCR difference is −3.3 points (95% CI [−6.3, −0.3]) and is not significant after correction (adjusted p = .100); relative to Neutral-V2, the difference is −7.9 points (95% CI [−11.4, −4.6]; adjusted p = 3.59×10⁻5). Neutral-V2 exceeds Neutral-V1 by 4.6 points despite identical token length (95% CI [2.2, 7.1]; adjusted p = .00146). Confidence intervals are unadjusted paired-bootstrap intervals, whereas the quoted p values are Holm-adjusted. The large Exp2–Exp3 difference is therefore reproduced primarily by interleaved assistant-turn and contextual structure. Self-history content may provide a smaller additional benefit, but its estimated magnitude and significance depend on the neutral wording."
    ], field_bank)

    # Robustness wording.
    set_paragraph(P[90], ["Robustness on a non-overlapping 500-example subset"], field_bank)
    set_paragraph(P[92], [
        "We repeated the main protocol evaluation on a second, non-overlapping 500-example subset from later validation rows. The qualitative ordering was consistent with the primary tables: Qwen and Phi remained most vulnerable under sequential-final exposure, while Gemma remained less susceptible overall and did not show the same ordering between the sequential protocols. This replication shows that the primary pattern is not unique to the first contiguous 500 examples, but it should not be interpreted as an estimate of sampling variance across many independently drawn subsets."
    ], field_bank)

    # Log-probability and trajectory wording; table numbers after new Table 5.
    set_paragraph(P[97], [
        "We next examine whether false consensus changes the model’s option-level preference. Table 6 reports the average log-probabilities of the target distractor and the correct answer before and after exposure, together with their shifts. For Qwen and Phi, Exp2 is associated with a substantial increase in target-distractor log-probability and a decrease in correct-answer log-probability. Gemma shows much smaller shifts overall. Exp3 is associated with smaller movement toward the distractor for Qwen and Phi, but this contrast should be interpreted at the protocol level rather than as evidence that intermediate answering alone causes the difference."
    ], field_bank)
    set_paragraph(P[98], ["Table 6. Log-probability metrics from the initial no-interference state"], field_bank)
    set_paragraph(P[103], ["Table 7. Model susceptibility ranking"], field_bank)
    set_paragraph(P[106], [
        "Stepwise trajectories reveal different dynamics in Exp3 and Exp4. For Qwen and Phi, Exp4 generally shows lower conformity at Step 5 than at Step 1, although the trajectories are not strictly monotonic. Exp4 nevertheless remains above Exp3 for most evaluated settings. The Exp3 trajectory for Qwen stays low throughout, whereas the Gemma trajectory begins higher and remains comparatively stable. These trajectories describe how the complete protocols evolve over turns; they do not independently identify intermediate answering as the causal source of the difference."
    ], field_bank)
    set_paragraph(P[107], ["Table 8. Conformity-rate trajectories for Exp3 and Exp4"], field_bank)

    # Make the wide log-probability table compact without changing its font family.
    log_table = T[8]
    short_headers = ["Data", "Model", "Prot.", "N", "D before", "D after", "ΔD", "C before", "C after", "ΔC"]
    set_table_widths(log_table, [0.52, 0.60, 0.46, 0.42, 0.66, 0.66, 0.53, 0.66, 0.66, 0.53])
    for r_idx, row in enumerate(log_table.rows):
        prevent_row_split(row, header=(r_idx == 0))
        for c_idx, cell in enumerate(row.cells):
            txt = short_headers[c_idx] if r_idx == 0 else " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            set_cell_text(cell, txt, bold=(r_idx == 0), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_margins(cell, top=35, start=35, bottom=35, end=35)

    # Valid/invalid output table after Figure 6 image and before social-framing results.
    fig6_table = T[12]
    validity_heading = make_paragraph_after(fig6_table, P[84], ["Output validity across the main protocols"], field_bank)
    validity_intro = make_paragraph_after(validity_heading, P[85], [
        "Table 9 reports the number of valid and invalid outputs among the 500 attempted items in each primary model–dataset–protocol cell. Across-protocol inference should use common-valid item identifiers when these counts differ."
    ], field_bank)
    validity_caption = make_paragraph_after(validity_intro, P[86], ["Table 9. Valid and invalid outputs in the primary evaluation"], field_bank)
    validity_data = [["Dataset", "Model", "Protocol", "Valid", "Invalid"]]
    for dataset, model, vals in [
        ("CSQA", "Qwen2.5-3B", [500, 500, 500, 500]),
        ("CSQA", "Phi-3.5-mini", [497, 497, 497, 497]),
        ("CSQA", "Gemma2-2B", [500, 500, 500, 500]),
        ("MMLU", "Qwen2.5-3B", [500, 500, 500, 500]),
        ("MMLU", "Phi-3.5-mini", [494, 496, 489, 482]),
        ("MMLU", "Gemma2-2B", [499, 499, 499, 499]),
    ]:
        for exp, valid in enumerate(vals, start=1):
            validity_data.append([dataset, model, f"Exp{exp}", str(valid), str(500 - valid)])
    validity_table = insert_table_after(validity_caption, doc, validity_data, [0.75, 1.45, 0.85, 0.65, 0.65], font_size=9)
    validity_note = make_paragraph_after(validity_table, P[87], [
        "Note. Invalid denotes an attempted item for which the required initial or final answer could not be parsed into one valid option label."
    ], field_bank)
    for run in validity_note.runs:
        run.italic = True
        run.font.size = Pt(10)

    # Social framing results and synthesis.
    set_paragraph(P[110], ["Exploratory social-framing results for Gemma"], field_bank)
    set_paragraph(P[111], [
        "The Gemma results in the main protocols indicate low overall susceptibility under the short “Agent j: dᵢ” format. The auxiliary CommonsenseQA comparison suggests that responses differ when the same wrong recommendation is presented with explicit numbered-model framing rather than unlabelled repeated wording. The mixed-support condition produces lower distractor adoption and higher correct-answer selection than the unanimously wrong labelled condition. However, the labelled-versus-unlabelled contrast includes unavoidable wording differences, and the mixed condition changes both group composition and opinion correctness. These findings are therefore treated as exploratory evidence of sensitivity to social framing and support composition, not as an isolated or general causal effect of labels."
    ], field_bank)
    implications = [
        "These results have four implications for LLM reliability under misleading social or user-provided signals",
        FieldRef(113, 0),
        ". First, identical wrong-peer evidence can produce markedly different outcomes when it is organized through different interaction protocols. Second, the Qwen–CommonsenseQA ablation shows that longer interleaved neutral assistant content, rather than assistant-role markers alone, reproduces most of the original sequential-protocol difference. Third, self-history content may contribute additional resistance, but its estimated effect is smaller and sensitive to neutral wording. Fourth, option-level monitoring complements final-answer metrics by revealing shifts toward a misleading alternative even when the final choice does not change. The mechanism findings should be applied only to the evaluated Qwen–CommonsenseQA setting until replicated across models and datasets."
    ]
    set_paragraph(P[113], implications, field_bank)
    limitations = [
        "This study has several limitations. First, the peer opinions are controlled textual signals rather than outputs from independently executed autonomous agents; this improves experimental control but does not reproduce dynamic multi-agent interaction",
        FieldRef(115, 0),
        ". Second, the main comparison covers three relatively small open-source instruction-tuned checkpoints and two multiple-choice benchmarks, so the findings should not be generalized to closed models, larger scales, or open-ended dialogue. Third, the mechanism ablation is limited to Qwen2.5-3B-Instruct on CommonsenseQA. The significant difference between two token-length-matched neutral phrasings prevents a unique attribution to self-history content, and the Answer-History versus Answer+Confidence comparison is not itself token-length matched. Fourth, the robustness evaluation uses one additional non-overlapping 500-example subset rather than many independently sampled subsets. Fifth, the number of wrong peer signals is fixed at five, and the primary study does not test correct peer consensus or rationale-bearing peer messages. Finally, the Gemma social-framing analysis is exploratory and changes more than one feature in the mixed-support condition. Future work should replicate the ablation across models and datasets, systematically vary neutral wording and peer count, and use label-only and rationale-matched controls."
    ]
    set_paragraph(P[115], limitations, field_bank)

    # Conclusions.
    set_paragraph(P[117], [
        "This study examines protocol-dependent conformity under controlled wrong-peer evidence. Holding the question, target distractor, peer count, peer content, and peer order fixed, the four-protocol comparison shows large behavioral and log-probability differences among the evaluated models. Qwen and Phi exhibit the highest harmful conformity under Sequential-Final and substantially lower conformity under Sequential-Stepwise, whereas Gemma is less susceptible overall and shows a different ordering. These findings establish a model-dependent protocol difference but do not, by themselves, isolate intermediate answering as its cause."
    ], field_bank)
    set_paragraph(P[118], [
        "The Qwen–CommonsenseQA mechanism ablation substantially changes the interpretation of this difference. Two independently worded, 13-token neutral assistant controls reproduce most of the reduction in harmful conformity, while a three-token acknowledgment control does not. Answer-plus-confidence history yields a smaller additional benefit whose magnitude and statistical significance depend on the neutral wording. The evidence therefore points primarily to interleaved assistant-turn and contextual structure, with a smaller and control-dependent contribution from self-history content."
    ], field_bank)
    set_paragraph(P[119], [
        "Option-level log-probability analyses further show that wrong-peer evidence can shift relative preference toward the target distractor, particularly for Qwen and Phi under Sequential-Final. The target-probability and Gemma social-framing analyses provide supplementary evidence about initial distractor plausibility and sensitivity to presentation, but these auxiliary findings should not be interpreted as broad causal effects."
    ], field_bank)
    set_paragraph(P[120], [
        "Together, the results show that interaction design is a reliability factor in the evaluated open-source multiple-choice settings. Robust systems should monitor both final answers and preference shifts, and should not assume that stepwise interaction is universally protective. Replication across additional model scales, families, datasets, and open-ended settings is required before extending the mechanism conclusions beyond the present scope."
    ], field_bank)

    # Keep the data statement compact enough to avoid a one-line spill.
    set_paragraph(P[127], [
        "Code, prompts, evaluation scripts, and processed results are available at https://github.com/xuexucheng/Social-Conformity-in-Large-Language-Models. CommonsenseQA and MMLU were obtained from public Hugging Face datasets and remain subject to their original licenses. No human-subject data were collected."
    ], field_bank)

    # Update the AI-assistance disclosure to match this revision workflow.
    set_paragraph(P[130], [
        "Large language models were used as experimental subjects in this study, including Qwen2.5-3B-Instruct, Phi-3.5-mini-instruct, and Gemma2-2B-it. These models were used only for the experimental procedures described in the Experimental section. During manuscript preparation and revision, ChatGPT and Codex were used to assist with language polishing, consistency checking, document organization, and drafting text based on analyses supplied and verified by the authors. The authors reviewed, edited, and verified all AI-assisted text, statistical values, and interpretations. No AI tool was listed as an author, and the authors take full responsibility for the accuracy and integrity of the manuscript."
    ], field_bank)

    # Global pagination and layout controls.
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if text.startswith("Figure ") or text.startswith("Table "):
            set_keep(p, with_next=True, together=True)
        if text.isupper() or text in {
            "Related work and positioning", "Study design and main findings", "Main contributions",
            "Tasks and datasets", "Target-distractor selection", "Models",
            "Synthetic peer-opinion construction", "Interaction protocols",
            "Mechanism-oriented ablation of the sequential protocol difference",
            "Exploratory social-framing analysis", "Metrics", "Generation, output format, and parsing",
            "Implementation and reproducibility details",
            "Statistical and robustness analysis", "Main behavioral results",
            "Mechanism ablation identifies a dominant contextual-structure effect",
            "Output validity across the main protocols", "Implications", "Limitations and future work",
            "AI and AI-assisted tools statement",
        }:
            set_keep(p, with_next=True, together=True)

    for table in doc.tables:
        for r_idx, row in enumerate(table.rows):
            prevent_row_split(row, header=(r_idx == 0 and len(table.rows) > 1))

    # Normalize core manuscript fonts without altering heading sizes or paragraph geometry.
    for p in doc.paragraphs:
        for run in p.runs:
            if run.text:
                run.font.name = "Times New Roman"
                rpr = run._element.get_or_add_rPr()
                fonts = rpr.get_or_add_rFonts()
                fonts.set(qn("w:ascii"), "Times New Roman")
                fonts.set(qn("w:hAnsi"), "Times New Roman")

    # Save clean copy. Zotero fields outside revised paragraphs remain untouched;
    # revised citation-bearing paragraphs reinsert their original field XML.
    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)}")
    print(f"FIELDS={sum(len(p._p.xpath('.//w:instrText')) for p in doc.paragraphs)}")


if __name__ == "__main__":
    main()
